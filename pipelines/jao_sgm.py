#!/usr/bin/env python3
"""WP-JAO — JAO Core Static Grid Model (SGM) downloader + parser.

The Core SGM is the only market-published, element-level network model for the
Core flow-based region (DE/FR/PL/AT/BE/NL/CZ/SI/HR/HU/RO/SK/... TSOs). It ships as
a ~6-monthly Excel workbook of lines, tie-lines and transformers/PSTs with real
electrical parameters (R, X, B), current ratings (Imax) and voltage levels — the
physical network our DC power-flow engine calibrates against (POWERFLOW_PLAN §2.3b).

This script:
  1. discovers the latest release ZIP on https://www.jao.eu/static-grid-model
     (falls back to a known-good URL if the page layout changes),
  2. downloads + caches it immutably under pipelines/.cache/sgm/,
  3. parses the Lines / Tielines / Transformers sheets of the "*_External.xlsx"
     workbook into a normalized bus/branch graph, and
  4. writes data/networks/core_sgm.json.

Output schema:
  {
    "meta":  {source, release, url, sha256, fetched_at, units, counts, ...},
    "buses": [ {id, name, voltage_kv} ],                 # id = "<substation>@<kV>"
    "branches": [ {from, to, type, r, x, b, imax_a, voltage_kv, tso, name, eic} ]
  }
  type ∈ {line, tieline, transformer, pst}

Units: r,x in ohms; b in microsiemens (µS); imax_a in amperes; voltage_kv in kV.

Usage:
  python3 pipelines/jao_sgm.py                 # discover → download → parse → write
  python3 pipelines/jao_sgm.py --url <zip>     # pin a specific release ZIP
  python3 pipelines/jao_sgm.py --zip <path>    # parse an already-downloaded ZIP
  python3 pipelines/jao_sgm.py --force         # re-download even if cached
"""
import argparse
import hashlib
import json
import os
import re
import sys
import zipfile
from datetime import datetime, timezone

import openpyxl
import requests

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, "pipelines", ".cache", "sgm")
OUT = os.path.join(ROOT, "data", "networks", "core_sgm.json")
SGM_PAGE = "https://www.jao.eu/static-grid-model"
JAO_BASE = "https://www.jao.eu"
USER_AGENT = "NeuraEnergy-OpportunityMap/WP-JAO (research; contact karimcy@gmail.com)"

# Known-good latest release as of 2026-07 (9th release, dated 2026-03-01); used
# as a fallback if page discovery fails.
FALLBACK_URL = (JAO_BASE +
    "/sites/default/files/2026-06/202603_Core%20Static%20Grid%20Model_9th%20release.zip")
FALLBACK_LABEL = "9th release (2026-03-01)"


# --- discovery ---------------------------------------------------------------
def discover_latest_url():
    """Scrape the SGM page and return (url, label) for the highest release ordinal.

    Picks by the ordinal in "<n>th release" so 'outdated_'-prefixed and reissued
    files sort correctly. Falls back to the pinned known-good URL on any failure.
    """
    try:
        r = requests.get(SGM_PAGE, headers={"User-Agent": USER_AGENT}, timeout=30)
        r.raise_for_status()
        links = re.findall(r'href="([^"]*[Ss]tatic[^"]*\.zip[^"]*)"', r.text)
        best = None  # (ordinal, url)
        for href in links:
            dec = requests.utils.unquote(href)
            m = re.search(r"(\d+)\s*(?:st|nd|rd|th)\s+release", dec, re.I)
            if not m:
                continue
            ordn = int(m.group(1))
            if best is None or ordn > best[0]:
                best = (ordn, href)
        if best:
            url = best[1] if best[1].startswith("http") else JAO_BASE + best[1]
            label = f"{best[0]}th release"
            return url, label
    except Exception as e:
        print(f"  ! discovery failed ({e}); using fallback URL", file=sys.stderr)
    return FALLBACK_URL, FALLBACK_LABEL


# --- download ----------------------------------------------------------------
def download_sgm(url, force=False):
    """Download the release ZIP to the immutable cache. Returns (zip_path, sha256)."""
    os.makedirs(CACHE, exist_ok=True)
    fname = requests.utils.unquote(os.path.basename(url))
    fname = re.sub(r"[^\w.\- ]", "_", fname)
    path = os.path.join(CACHE, fname)
    if os.path.exists(path) and not force:
        data = open(path, "rb").read()
        return path, hashlib.sha256(data).hexdigest()
    print(f"  downloading {url}")
    r = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=180)
    r.raise_for_status()
    tmp = path + ".tmp"
    with open(tmp, "wb") as f:
        f.write(r.content)
    os.replace(tmp, path)
    return path, hashlib.sha256(r.content).hexdigest()


def extract_xlsx(zip_path):
    """Return the path to the extracted '*_External.xlsx' workbook."""
    outdir = os.path.join(CACHE, "extract")
    with zipfile.ZipFile(zip_path) as z:
        z.extractall(outdir)
    for root, _, files in os.walk(outdir):
        for fn in files:
            if fn.lower().endswith("external.xlsx"):
                return os.path.join(root, fn)
    # fall back to any .xlsx
    for root, _, files in os.walk(outdir):
        for fn in files:
            if fn.lower().endswith(".xlsx"):
                return os.path.join(root, fn)
    raise FileNotFoundError("no .xlsx found inside SGM ZIP")


# --- parsing helpers ---------------------------------------------------------
def _num(v):
    """Coerce a cell to float, or None. Handles blanks, spaces, commas."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _clean(v):
    return None if v is None else str(v).strip() or None


def _rows(ws):
    """Data rows (skip the two header rows), each as a value tuple."""
    return list(ws.iter_rows(min_row=3, values_only=True))


def _pick_imax_line(row):
    """Line/tieline Imax: prefer Fixed, else max seasonal period, else DLRmax."""
    fixed = _num(row[12])
    if fixed:
        return fixed
    seasonal = [x for x in (_num(row[i]) for i in range(6, 12)) if x]
    if seasonal:
        return max(seasonal)
    return _num(row[14])  # DLRmax


def _pick_imax_tr(row):
    """Transformer Imax: prefer Fixed, else Max, else Min."""
    return _num(row[5]) or _num(row[4]) or _num(row[3])


def _bus_id(name, kv):
    return f"{name}@{int(kv)}" if kv and float(kv).is_integer() else f"{name}@{kv}"


def parse_sgm(xlsx_path):
    """Parse the SGM workbook into (buses, branches, sheet_counts)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    buses = {}   # id -> {id, name, voltage_kv}
    branches = []
    counts = {}

    def add_bus(name, kv):
        if not name or kv is None:
            return None
        bid = _bus_id(name, kv)
        if bid not in buses:
            buses[bid] = {"id": bid, "name": name, "voltage_kv": kv}
        return bid

    # --- Lines & Tielines share a layout ---
    for sheet, btype in (("Lines", "line"), ("Tielines", "tieline")):
        if sheet not in wb.sheetnames:
            continue
        n = 0
        for row in _rows(wb[sheet]):
            name = _clean(row[0])
            if not name:
                continue
            eic = _clean(row[1])
            tso = _clean(row[2])
            sub1, sub2 = _clean(row[3]), _clean(row[4])
            kv = _num(row[5])
            r, x, b = _num(row[15]), _num(row[16]), _num(row[17])
            imax = _pick_imax_line(row)
            fbus = add_bus(sub1, kv)
            tbus = add_bus(sub2, kv)
            branches.append({
                "from": fbus, "to": tbus, "type": btype,
                "r": r, "x": x, "b": b, "imax_a": imax, "voltage_kv": kv,
                "tso": tso, "name": name, "eic": eic,
                "length_km": _num(row[18]),
            })
            n += 1
        counts[sheet] = n

    # --- Transformers (incl. PSTs) ---
    if "Transformers" in wb.sheetnames:
        n = 0
        for row in _rows(wb["Transformers"]):
            name = _clean(row[0])
            if not name:
                continue
            eic = _clean(row[1])
            tso = _clean(row[2])
            v_pri, v_sec = _num(row[6]), _num(row[7])
            r, x, b = _num(row[8]), _num(row[9]), _num(row[10])
            imax = _pick_imax_tr(row)
            # A phase-shifting transformer (PST) has quadrature/angle regulation:
            # Sym/Asym type set, non-zero angle-regulation δu, or a Theta angle.
            # An OLTC transformer with only in-phase Phase-Regulation δu (voltage
            # tapping) is NOT a PST and stays type "transformer".
            theta, angle_reg = _num(row[13]), _num(row[16])
            sym_asym = _clean(row[14])
            is_pst = bool(sym_asym) or bool(angle_reg) or bool(theta)
            btype = "pst" if is_pst else "transformer"
            # transformer connects the two voltage buses at the same location
            fbus = add_bus(name, v_pri)
            tbus = add_bus(name, v_sec)
            br = {
                "from": fbus, "to": tbus, "type": btype,
                "r": r, "x": x, "b": b, "imax_a": imax,
                "voltage_kv": v_pri, "voltage_secondary_kv": v_sec,
                "tso": tso, "name": name, "eic": eic,
            }
            if is_pst:  # keep the phase-shifter parameters for DC-PF modelling
                br["pst"] = {
                    "sym_asym": sym_asym, "angle_reg_pct": angle_reg,
                    "theta_deg": theta, "taps": _clean(row[12]),
                    "phase_reg_pct": _num(row[15]),
                }
            branches.append(br)
            n += 1
        counts["Transformers"] = n

    # supplementary: count Remedial Actions (PSTs listed as RAs) for provenance
    if "Remedial Actions" in wb.sheetnames:
        counts["Remedial Actions"] = sum(
            1 for row in _rows(wb["Remedial Actions"]) if _clean(row[0]))

    return list(buses.values()), branches, counts


# --- write -------------------------------------------------------------------
def write_output(buses, branches, counts, provenance):
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    # per-type / per-TSO tallies for the meta block
    by_type, by_tso = {}, {}
    for br in branches:
        by_type[br["type"]] = by_type.get(br["type"], 0) + 1
        by_tso.setdefault(br["tso"] or "?", {})
        by_tso[br["tso"] or "?"][br["type"]] = by_tso[br["tso"] or "?"].get(br["type"], 0) + 1
    doc = {
        "meta": {
            **provenance,
            "units": {"r": "ohm", "x": "ohm", "b": "microsiemens",
                      "imax_a": "ampere", "voltage_kv": "kV"},
            "sheet_counts": counts,
            "bus_count": len(buses),
            "branch_count": len(branches),
            "branches_by_type": by_type,
            "branches_by_tso": by_tso,
        },
        "buses": sorted(buses, key=lambda b: (b["name"] or "", b["voltage_kv"] or 0)),
        "branches": branches,
    }
    tmp = OUT + ".tmp"
    with open(tmp, "w") as f:
        json.dump(doc, f, indent=1, ensure_ascii=False)
    os.replace(tmp, OUT)
    return doc


# --- CLI ---------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description="Download + parse the JAO Core Static Grid Model")
    ap.add_argument("--url", help="release ZIP URL (default: auto-discover latest)")
    ap.add_argument("--zip", help="use an already-downloaded ZIP instead of downloading")
    ap.add_argument("--force", action="store_true", help="re-download even if cached")
    args = ap.parse_args(argv)

    if args.zip:
        zip_path = args.zip
        sha = hashlib.sha256(open(zip_path, "rb").read()).hexdigest()
        url, label = "(local)", os.path.basename(zip_path)
    else:
        if args.url:
            url, label = args.url, "user-specified"
        else:
            url, label = discover_latest_url()
            print(f"  latest release: {label}\n  {url}")
        zip_path, sha = download_sgm(url, force=args.force)

    print(f"  ZIP: {zip_path}  sha256={sha[:16]}…")
    xlsx = extract_xlsx(zip_path)
    print(f"  workbook: {os.path.basename(xlsx)}")

    buses, branches, counts = parse_sgm(xlsx)
    provenance = {
        "source": "JAO Core Static Grid Model",
        "release": label,
        "url": url,
        "zip_sha256": sha,
        "workbook": os.path.basename(xlsx),
        "fetched_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "licence": "JAO publication (free, no registration). Attribution: JAO SAO.",
    }
    doc = write_output(buses, branches, counts, provenance)

    m = doc["meta"]
    print(f"\n  buses: {m['bus_count']}   branches: {m['branch_count']}")
    print(f"  by type: {m['branches_by_type']}")
    print(f"  sheets:  {m['sheet_counts']}")
    print(f"  wrote {OUT}")


if __name__ == "__main__":
    main()
